"""에어코리아 최종확정 측정자료 → 측정소별 3년 평균 집계 (세션559)

## 왜 python 인가
이 저장소에는 xlsx 를 읽는 선례가 0건이다(grep 확인). Node 에 새 의존성을 들이는 대신
이미 깔려 있는 openpyxl(3.1.5)로 집계해 JSON 을 만들고, Node 쪽 반영 스크립트가 그 JSON 을
DB 에 넣는 2단 구조를 쓴다. 파일 하나가 300~400MB 라 read_only 스트리밍이 필수다.

## 입력
data.go.kr 15122830 "한국환경공단_에어코리아_최종확정 측정자료" 연도별 XLSX.
신청 불필요·무료. 파일 다운로드는 사람이 하거나 브라우저 자동화로 받는다(연 1회).

## ⚠️ 연도마다 형식이 다르다 (2026-09-22 실측)
| 연도 | 시트 이름     | 측정일시 형식      |
|------|--------------|-------------------|
| 2022 | "1월"         | 2022010101 (숫자)  |
| 2023 | "2023년 1월"  | "2023-01-01-01" (문자열) |
| 2024 | "2024년 1월 " | 2024010101 (숫자, 시트명 뒤 공백) |

그래서 **시트를 이름으로 찾지 않는다** — wb.sheetnames 를 통째로 순회한다(어차피 12개 전부 읽는다).
컬럼도 위치를 박지 않고 **헤더 행에서 인덱스를 찾는다**. 형식이 또 바뀌어도 견딘다.

측정일시는 **쓰지 않는다** — 연 평균이라 시각이 필요 없다. 그래서 숫자/문자열 차이가 무해하다.

## 사용
    python scripts/air-annual-aggregate.py --files F:/tmp/airkorea_2022.xlsx F:/tmp/airkorea_2023.xlsx ... \\
        --years 2022,2023,2024 --out artifacts/air-annual.json
"""

import argparse
import io
import json
import os
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

try:
    import openpyxl
except ImportError:
    print("openpyxl 이 없습니다. `pip install openpyxl` 후 다시 실행하세요.")
    sys.exit(1)

# 헤더에서 찾을 컬럼 — 부분 일치로 찾는다(연도별 표기가 미세하게 다를 수 있다)
COL_PATTERNS = {
    "station": ["측정소명"],
    "code": ["측정소코드"],
    "addr": ["주소"],
    "pm10": ["미세먼지", "PM10"],   # ⚠️ "미세먼지(PM10)" 과 "초미세먼지(PM25)" 가 둘 다 "미세먼지" 를
    "pm25": ["초미세먼지", "PM25"],  #    포함하므로, pm25 를 **먼저** 찾아 그 인덱스를 제외한다
    "o3": ["오존", "O3"],
}


def find_columns(header):
    """헤더 행에서 컬럼 인덱스를 찾는다. 못 찾으면 None.

    ⚠️ '미세먼지(PM10)' 과 '초미세먼지(PM25)' 는 부분 일치가 겹친다.
    초미세먼지를 먼저 확정하고 그 인덱스를 미세먼지 후보에서 뺀다.
    """
    cells = [str(c) if c is not None else "" for c in header]
    idx = {}

    def search(keys, exclude=()):
        for i, c in enumerate(cells):
            if i in exclude:
                continue
            for k in keys:
                if k in c:
                    return i
        return None

    idx["pm25"] = search(COL_PATTERNS["pm25"])
    taken = {idx["pm25"]} if idx["pm25"] is not None else set()
    idx["pm10"] = search(COL_PATTERNS["pm10"], exclude=taken)
    for key in ("station", "code", "addr", "o3"):
        idx[key] = search(COL_PATTERNS[key])
    return idx


def aggregate(paths):
    """{station: {pm25합, pm25수, pm10합, pm10수, o3합, o3수, code, addr}}"""
    acc = {}
    for path in paths:
        if not os.path.exists(path):
            print(f"  ⚠️ 파일 없음, 건너뜀: {path}")
            continue
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        total = 0
        for sheet in wb.sheetnames:  # 이름으로 찾지 않는다 — 전부 읽는다
            ws = wb[sheet]
            # 헤더를 못 찾으면 그 시트는 통째로 건너뛴다. idx 가 None 인 채로 아래 루프에
            # 들어가면 조용히 엉뚱한 컬럼을 읽거나 터진다 — 그래서 시트 단위로 먼저 확정한다.
            idx = None
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n == 0:
                    found = find_columns(row)
                    if found.get("station") is None or found.get("pm25") is None:
                        print(f"  ⚠️ {os.path.basename(path)} / {sheet!r}: 헤더에서 컬럼을 못 찾음 — 건너뜀")
                        break
                    idx = found
                    continue
                if idx is None:
                    break
                total += 1
                st = row[idx["station"]]
                if not st:
                    continue
                st = str(st).strip()
                a = acc.setdefault(st, {"s25": 0.0, "n25": 0, "s10": 0.0, "n10": 0,
                                        "s3": 0.0, "n3": 0, "code": None, "addr": None})
                # 코드·주소는 처음 본 값을 쓴다(같은 측정소면 동일)
                if a["code"] is None and idx.get("code") is not None:
                    v = row[idx["code"]]
                    if v is not None:
                        a["code"] = str(v).strip()
                if a["addr"] is None and idx.get("addr") is not None:
                    v = row[idx["addr"]]
                    if v:
                        a["addr"] = str(v).strip()
                # 빈 칸은 점검·통신장애로 인한 미수신이다 — 0 으로 세지 않고 제외한다
                for key, sk, nk in (("pm25", "s25", "n25"), ("pm10", "s10", "n10"), ("o3", "s3", "n3")):
                    i = idx.get(key)
                    if i is None:
                        continue
                    v = row[i]
                    if isinstance(v, (int, float)):
                        a[sk] += float(v)
                        a[nk] += 1
        wb.close()
        print(f"  {os.path.basename(path)}: {total:,}행")
    return acc


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--files", nargs="+", required=True, help="연도별 XLSX 경로")
    ap.add_argument("--years", required=True, help='집계 연도 표기 (예: "2022,2023,2024")')
    ap.add_argument("--out", required=True, help="출력 JSON 경로")
    args = ap.parse_args()

    print(f"집계 시작 — 파일 {len(args.files)}개")
    acc = aggregate(args.files)

    out = []
    for st, a in sorted(acc.items()):
        if a["n25"] == 0 and a["n10"] == 0:
            continue  # 측정값이 하나도 없는 측정소는 제외
        out.append({
            "station_name": st,
            "station_code": a["code"],
            "pm25": round(a["s25"] / a["n25"], 2) if a["n25"] else None,
            "pm10": round(a["s10"] / a["n10"], 2) if a["n10"] else None,
            "o3": round(a["s3"] / a["n3"], 5) if a["n3"] else None,
            "years": args.years,
            "sample_hours": a["n25"] or a["n10"],
            "address": a["addr"],
        })

    os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
    with io.open(args.out, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)

    pm25 = sorted(r["pm25"] for r in out if r["pm25"] is not None)
    print(f"\n측정소 {len(out)}곳 → {args.out}")
    if pm25:
        q = lambda p: pm25[int(len(pm25) * p)]
        print(f"pm25 최소={pm25[0]} 25%={q(.25)} 중앙={q(.5)} 75%={q(.75)} 최대={pm25[-1]}")
    thin = [r for r in out if (r["sample_hours"] or 0) < 8760]  # 1년치 미만
    if thin:
        print(f"⚠️ 표본 8,760시간(1년) 미만 측정소 {len(thin)}곳 — 결측이 많거나 중간 신설/폐지")


if __name__ == "__main__":
    main()
